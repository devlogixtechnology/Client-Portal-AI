"""
Excel Financial Statement Extractor using openpyxl.
Extracts structured sheets (Balance Sheet, Income Statement) and text disclosures/notes.
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from src.extractors.base import BaseExtractor
from src.models import ExtractedRawData


class ExcelExtractor(BaseExtractor):
    """Extracts tables, narrative disclosures, and metadata from Excel financial workbooks."""

    def extract(self, file_path: Path) -> ExtractedRawData:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {file_path}")

        all_tables: List[List[List[Any]]] = []
        narrative_sections: List[Dict[str, str]] = []
        full_text_chunks: List[str] = []

        company_name: Optional[str] = None
        fiscal_period: Optional[str] = None
        currency: str = "USD"
        unit_multiplier: int = 1

        wb = openpyxl.load_workbook(filename=str(file_path), data_only=True)

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            sheet_rows: List[List[str]] = []
            is_notes_sheet = bool(re.search(r"notes|disclosure|footnote|policy", sheetname, re.IGNORECASE))
            sheet_text_lines: List[str] = []

            for row in ws.iter_rows(values_only=True):
                # Filter out completely empty rows
                if not any(c is not None and str(c).strip() for c in row):
                    continue

                cleaned_row = [str(c).strip() if c is not None else "" for c in row]
                row_str = " ".join([c for c in cleaned_row if c])
                sheet_text_lines.append(row_str)

                # Look for metadata in first few non-empty rows
                if not company_name and len(cleaned_row) > 0 and cleaned_row[0]:
                    first_cell = cleaned_row[0]
                    if not any(kw in first_cell.lower() for kw in ["balance sheet", "income statement", "period", "date", "year"]):
                        company_name = first_cell

                if currency == "USD":
                    currency = self.detect_currency(row_str)
                if unit_multiplier == 1:
                    unit_mult, _ = self.detect_unit_multiplier(row_str)
                    if unit_mult != 1:
                        unit_multiplier = unit_mult
                if not fiscal_period:
                    detected_p = self.detect_fiscal_period(row_str)
                    if detected_p != "FY":
                        fiscal_period = detected_p

                # Check if row is purely narrative (e.g. single cell with long text or Note label)
                non_empty_cells = [c for c in cleaned_row if c]
                if len(non_empty_cells) == 1 and len(non_empty_cells[0]) > 40:
                    # Likely a narrative note line inside a sheet
                    if re.search(r"^Note\s+\d+|Accounting Policies|Commitments", non_empty_cells[0], re.IGNORECASE):
                        narrative_sections.append({
                            "sheet": sheetname,
                            "title": non_empty_cells[0][:80],
                            "text": non_empty_cells[0]
                        })
                else:
                    sheet_rows.append(cleaned_row)

            full_text_chunks.append(f"--- SHEET: {sheetname} ---\n" + "\n".join(sheet_text_lines))

            if is_notes_sheet:
                narrative_sections.append({
                    "sheet": sheetname,
                    "title": f"Disclosures & Notes ({sheetname})",
                    "text": "\n".join(sheet_text_lines)
                })
            elif len(sheet_rows) >= 2:
                all_tables.append(sheet_rows)

        # Fallback company name from file name
        if not company_name:
            clean_stem = file_path.stem.replace("_", " ").replace("-", " ")
            company_name = clean_stem.split("FY")[0].strip()

        if not fiscal_period:
            fiscal_period = self.detect_fiscal_period(file_path.name)

        return ExtractedRawData(
            company_name=company_name,
            fiscal_period=fiscal_period,
            currency=currency,
            unit_multiplier=unit_multiplier,
            tables=all_tables,
            narrative_sections=narrative_sections,
            raw_text="\n\n".join(full_text_chunks),
            source_file=file_path.name,
            file_type="xlsx"
        )
